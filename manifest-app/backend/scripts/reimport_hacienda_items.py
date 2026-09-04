"""Reimporta el catálogo hacienda_items desde la fuente autoritativa de Hacienda.

Motivo: la tabla hacienda_items tenía ~1,230 códigos (23% del catálogo) con la
descripción de OTRO código (ej. 190531 mostraba "COOKING STOVES..." en vez de
"COOKIES") y 1,156 códigos faltaban por completo, frente al listado real
adjuntado por el usuario (Comodities.xlsx, hoja Sheet2: columnas Codigo y
Comodity, esta última un string de ancho fijo description[0:32]+unit[32:34]+
tariff[34:39]). Confirmado que hacienda_items solo alimenta el buscador del
editor (routes/catalogs.js) — no interviene en txtGenerator ni blValidation —
y que ningún B/L real tiene hacienda_item_code asignado todavía, así que
reemplazar la tabla completa es seguro.

Uso: python backend/scripts/reimport_hacienda_items.py <ruta_xlsx> <ruta_db>
"""
import sys
import sqlite3
import openpyxl


def normalizar_codigo(valor):
    if isinstance(valor, (int, float)):
        return str(int(valor))
    return str(valor).strip()


def parsear_fuente(ruta_xlsx):
    wb = openpyxl.load_workbook(ruta_xlsx, data_only=True)
    ws = wb['Sheet2']
    filas = ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True)
    catalogo = {}
    for codigo, comodity in filas:
        if codigo is None:
            continue
        code = normalizar_codigo(codigo)
        s = str(comodity) if comodity is not None else ''
        descripcion = s[0:32].strip()
        unidad = s[32:34].strip()
        tarifa = s[34:39].strip() or '00000'
        if not descripcion:
            continue
        # Códigos duplicados en la fuente: se conserva la última ocurrencia.
        catalogo[code] = (code, descripcion, unidad, tarifa, 1)
    return list(catalogo.values())


def main():
    if len(sys.argv) != 3:
        print('Uso: reimport_hacienda_items.py <xlsx> <db>')
        sys.exit(1)
    ruta_xlsx, ruta_db = sys.argv[1], sys.argv[2]

    filas = parsear_fuente(ruta_xlsx)
    print(f'Fuente parseada: {len(filas)} códigos únicos')

    con = sqlite3.connect(ruta_db)
    try:
        con.execute('BEGIN')
        antes = con.execute('SELECT COUNT(*) FROM hacienda_items').fetchone()[0]
        con.execute('DELETE FROM hacienda_items')
        con.executemany(
            'INSERT INTO hacienda_items (code, description, unit, tariff, taxable) VALUES (?,?,?,?,?)',
            filas,
        )
        despues = con.execute('SELECT COUNT(*) FROM hacienda_items').fetchone()[0]
        con.commit()
        print(f'hacienda_items: {antes} filas -> {despues} filas')
    except Exception:
        con.rollback()
        raise
    finally:
        con.close()


if __name__ == '__main__':
    main()
